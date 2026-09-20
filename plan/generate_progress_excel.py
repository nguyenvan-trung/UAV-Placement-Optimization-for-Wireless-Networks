"""Generate the Excel progress workbook. Run this file to reset the template."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT = Path(__file__).with_name("team_weekly_progress.xlsx")
MEMBERS = ("Nguyễn Văn Trung", "Hoàng Văn Trường", "Nguyễn Minh Hiếu")
ROLES = ("Chủ nhóm", "Thành viên", "Thành viên")
STATUSES = "Chưa bắt đầu,Đang thực hiện,Hoàn thành,Bị chặn"
WEEKS = (
    ("Học và xác định bài toán", "Tìm hiểu UAV placement, điều phối phạm vi", "Tìm hiểu GA và ứng dụng", "Tìm hiểu WOA, PSO và metric"),
    ("Đọc tài liệu và phân tích yêu cầu", "Nghiên cứu mô hình hệ thống và PSO", "Tổng hợp tài liệu GA và cách mã hóa", "Tổng hợp tài liệu WOA và cách đánh giá"),
    ("Thống nhất ý tưởng và logic code", "Chốt kiến trúc, interface và quy ước", "Đề xuất pipeline dữ liệu và module GA", "Đề xuất metric, biểu đồ và module WOA"),
    ("Dữ liệu và mô hình dùng chung", "Search space, encoding và interface", "Loader, validator và preprocessor", "User/UAV model và channel model"),
    ("Triển khai thuật toán phần 1", "Particle, initialization, velocity PSO", "Chromosome, initialization, selection GA", "Whale, initialization, coefficients WOA"),
    ("Triển khai thuật toán phần 2", "Hoàn thiện PSO, pbest/gbest", "Hoàn thiện crossover, mutation và GA", "Hoàn thiện encircling, spiral và WOA"),
    ("Tích hợp bài toán UAV", "Tích hợp scripts/run_optimizer.py, PSO và objective", "Tích hợp GA và kiểm tra dữ liệu", "Tích hợp WOA, metric và visualization"),
    ("Thử nghiệm và sửa lỗi", "Thiết kế kịch bản và kiểm tra công bằng", "Chạy GA, ghi log và kiểm tra lỗi", "Chạy đối chứng, tạo bảng và biểu đồ"),
    ("Thí nghiệm cuối và phân tích", "Tổng hợp, kiểm chứng và kết luận", "Phân tích GA, viết dữ liệu/phương pháp", "Phân tích kết quả, hoàn thiện biểu đồ"),
    ("Hoàn thiện báo cáo", "Ghép báo cáo, kiểm tra định dạng", "Rà soát lý thuyết và trích dẫn", "Rà soát kết quả, bảng, hình, phụ lục"),
)

NAVY, BLUE, WHITE = "1F4E78", "D9EAF7", "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def title(cell, size=15):
    cell.font = Font(bold=True, color=WHITE, size=size)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def header(cells):
    for cell in cells:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def status_validation(sheet, cells):
    validation = DataValidation(type="list", formula1=f'"{STATUSES}"')
    sheet.add_data_validation(validation)
    validation.add(cells)
    colors = {"Chưa bắt đầu": "D9E1F2", "Đang thực hiện": "FFF2CC", "Hoàn thành": "C6E0B4", "Bị chặn": "F4CCCC"}
    first = str(cells).split(":")[0]
    for label, color in colors.items():
        sheet.conditional_formatting.add(str(cells), FormulaRule(formula=[f'{first}="{label}"'], fill=PatternFill("solid", fgColor=color)))


wb = Workbook()
ws = wb.active
ws.title = "Tổng quan"
ws.merge_cells("A1:E1")
ws["A1"] = "THEO DÕI TIẾN ĐỘ ĐỀ TÀI UAV PLACEMENT"
title(ws["A1"])
ws.append(["Chủ nhóm", MEMBERS[0], "Thành viên", MEMBERS[1], MEMBERS[2]])
ws.append([])
ws.append(["Tuần", "Nội dung chính", "Trạng thái", "Ngày cập nhật", "Minh chứng / đường dẫn"])
header(ws[4])
for number, week in enumerate(WEEKS, 1):
    ws.append([number, week[0], "Chưa bắt đầu", "", ""])
    for cell in ws[number + 4]:
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)
status_validation(ws, "C5:C14")
ws.freeze_panes = "A5"
ws.auto_filter.ref = "A4:E14"
for column, width in zip("ABCDE", (9, 38, 20, 18, 45)):
    ws.column_dimensions[column].width = width

guide = wb.create_sheet("Hướng dẫn")
guide.merge_cells("A1:D1")
guide["A1"] = "HƯỚNG DẪN CẬP NHẬT"
title(guide["A1"])
steps = (
    "1. Mỗi người điền đúng dòng mang tên mình trong sheet của tuần hiện tại.",
    "2. Ghi công việc đã làm và dán link commit, Pull Request hoặc file kết quả.",
    "3. Chọn trạng thái trong danh sách có sẵn.",
    "4. Chủ nhóm điền tổng kết, khó khăn, kế hoạch tiếp theo và ngày xác nhận.",
    "5. Cập nhật dòng tương ứng trong sheet Tổng quan rồi commit file lên GitHub.",
    "6. GitHub cho phép xem/tải Excel; để cùng sửa trực tuyến nên mở bằng Excel Online hoặc Google Sheets.",
)
for row, text in enumerate(steps, 3):
    guide.cell(row, 1, text)
guide.column_dimensions["A"].width = 110

for number, week in enumerate(WEEKS, 1):
    sheet = wb.create_sheet(f"Tuần {number}")
    sheet.merge_cells("A1:F1")
    sheet["A1"] = f"TUẦN {number} — {week[0].upper()}"
    title(sheet["A1"], 14)
    sheet["A2"], sheet["B2"] = "Thời gian", "DD/MM/YYYY – DD/MM/YYYY"
    sheet["D2"], sheet["E2"] = "Trạng thái tuần", "Chưa bắt đầu"
    status_validation(sheet, "E2")
    sheet.append([])
    sheet.append(["Thành viên", "Vai trò", "Công việc dự kiến", "Công việc đã thực hiện", "Kết quả / minh chứng", "Trạng thái"])
    header(sheet[4])
    for index, member in enumerate(MEMBERS):
        sheet.append([member, ROLES[index], week[index + 1], "", "", "Chưa bắt đầu"])
        sheet.row_dimensions[5 + index].height = 55
        for cell in sheet[5 + index]:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    status_validation(sheet, "F5:F7")
    for row, label in enumerate(("Kết quả chung", "Khó khăn / vấn đề", "Kế hoạch tuần sau", "Nhận xét chủ nhóm", "Ngày xác nhận"), 9):
        sheet.cell(row, 1, label)
        sheet.cell(row, 1).font = Font(bold=True)
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=BLUE)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        for column in range(1, 7):
            sheet.cell(row, column).border = BORDER
            sheet.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = "A4:F7"
    for column, width in zip("ABCDEF", (23, 15, 38, 42, 42, 19)):
        sheet.column_dimensions[column].width = width

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUTPUT)
print(OUTPUT)
